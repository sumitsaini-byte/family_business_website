from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUTPUT_PATH = "/Users/sumit/Desktop/heeee copy 6/Sumit_Saini_Resume_Updated.docx"
PHOTO_PATH = "/private/tmp/pages_resume_extract/Data/pasted-movie-enhanced-42.png"


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        edge_data = kwargs.get(edge)
        if not edge_data:
            continue
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key, value in edge_data.items():
            element.set(qn(f"w:{key}"), str(value))


def add_run(paragraph, text, *, bold=False, italic=False, underline=False, size=10.5):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    return run


def add_section_heading(document, title):
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    add_run(p, title, bold=True, size=13)

    border_p = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    border_p.append(bottom)
    p._p.get_or_add_pPr().append(border_p)


def add_bullet(document, text, left=0.28, hanging=0.14):
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(left)
    p.paragraph_format.first_line_indent = Inches(-hanging)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_run(p, text, size=10.5)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.48)
section.bottom_margin = Inches(0.48)
section.left_margin = Inches(0.62)
section.right_margin = Inches(0.62)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(10.5)

header = doc.add_table(rows=1, cols=2)
header.autofit = False
header.columns[0].width = Inches(5.95)
header.columns[1].width = Inches(1.15)
for cell in header.rows[0].cells:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    set_cell_border(
        cell,
        top={"val": "nil"},
        left={"val": "nil"},
        bottom={"val": "nil"},
        right={"val": "nil"},
    )

left = header.cell(0, 0)
right = header.cell(0, 1)

p = left.paragraphs[0]
p.paragraph_format.space_after = Pt(2)
add_run(p, "SUMIT SAINI", bold=True, size=18)

contact_1 = left.add_paragraph()
contact_1.paragraph_format.space_after = Pt(0)
add_run(contact_1, "sumitsaini-byte · GitHub", underline=True, size=10.2)
add_run(contact_1, "    ", size=10.2)
add_run(contact_1, "linkedin.com/in/sumit-saini-25a340317", underline=True, size=10.2)

contact_2 = left.add_paragraph()
contact_2.paragraph_format.space_after = Pt(0)
add_run(contact_2, "sumit.sainisahb@gmail.com", underline=True, size=10.2)
add_run(contact_2, "    9654007110    Delhi Rohini, India.", size=10.2)

right_p = right.paragraphs[0]
right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
right_p.add_run().add_picture(PHOTO_PATH, width=Inches(1.08))

add_section_heading(doc, "Profile")
profile = doc.add_paragraph()
profile.paragraph_format.space_after = Pt(2)
profile.paragraph_format.line_spacing = 1.0
add_run(
    profile,
    "Motivated MCA student with strong skills in full-stack web development and real-world project experience. "
    "Passionate about applying technology in practical scenarios, especially in the finance domain. "
    "Quick learner, team-oriented, and open to international opportunities.",
)

add_section_heading(doc, "Education")

edu_1 = doc.add_paragraph()
edu_1.paragraph_format.space_after = Pt(0)
add_run(edu_1, "Master of Computer Applications (MCA), ", bold=True)
add_run(edu_1, "2024 - 2026", size=10.5)
edu_1.alignment = WD_ALIGN_PARAGRAPH.LEFT

edu_1b = doc.add_paragraph()
edu_1b.paragraph_format.space_after = Pt(1)
add_run(edu_1b, "Jagan Institute of Management Studies (JIMS), Rohini, Delhi", italic=True)

edu_2 = doc.add_paragraph()
edu_2.paragraph_format.space_after = Pt(0)
add_run(edu_2, "Bachelor of Computer Applications (BCA), Web Development, ", bold=True)
add_run(edu_2, "2021 - 2024", size=10.5)

edu_2b = doc.add_paragraph()
edu_2b.paragraph_format.space_after = Pt(0)
add_run(edu_2b, "Sirifort Institute of Management Studies, IP University, Sector 25, Rohini, Delhi", italic=True)

edu_2c = doc.add_paragraph()
edu_2c.paragraph_format.space_after = Pt(1)
add_run(edu_2c, "Percentage: 75%")

edu_3 = doc.add_paragraph()
edu_3.paragraph_format.space_after = Pt(0)
add_run(edu_3, "Senior Secondary (XII), Commerce, DAV Public School, Delhi    ", bold=True)
add_run(edu_3, "2020 - 2021")

edu_3b = doc.add_paragraph()
edu_3b.paragraph_format.space_after = Pt(2)
add_run(edu_3b, "CGPA: 7.25/10")

add_section_heading(doc, "Technical Skills")
skills = [
    ("Languages", "Java, C, JavaScript, Python, C++, PHP"),
    ("Web Development", "HTML5, CSS, React.js, Angular, Bootstrap"),
    ("Backend", "Node.js, Express.js, Spring Boot, REST APIs"),
    ("Database", "MySQL, MongoDB, SQL Server"),
    ("Cloud Basics", "Git, Docker"),
    ("Tools", "Postman, VS Code, IntelliJ, GitHub"),
    ("Concepts", "DSA, OOP"),
]
for label, value in skills:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_run(p, f"{label}: ", bold=True)
    add_run(p, value)

add_section_heading(doc, "Projects")

proj_1 = doc.add_paragraph()
proj_1.paragraph_format.space_after = Pt(0)
add_run(proj_1, "Hotel Management Website (Mar 2024 - May 2024)", bold=True)

proj_1_desc = doc.add_paragraph()
proj_1_desc.paragraph_format.space_after = Pt(0)
add_run(
    proj_1_desc,
    "A full-stack hotel management system for customer bookings and admin control.",
)
add_bullet(doc, "Built user login/register, room filtering, booking/cancellation system, and admin dashboard using PHP and MySQL.")
add_bullet(doc, "Designed a mobile-responsive UI using Bootstrap, focusing on usability and smooth session handling.")

proj_2 = doc.add_paragraph()
proj_2.paragraph_format.space_before = Pt(2)
proj_2.paragraph_format.space_after = Pt(0)
add_run(proj_2, "MERN E-commerce Website for Father's Business (Jul 2025 - Sep 2025)", bold=True)

proj_2_desc = doc.add_paragraph()
proj_2_desc.paragraph_format.space_after = Pt(0)
add_run(
    proj_2_desc,
    "Designed and developed a responsive e-commerce website for my father's men's fashion business.",
)
add_bullet(doc, "Built product catalog, customer authentication, cart, and order flow using MongoDB, Express.js, React.js, and Node.js.")
add_bullet(doc, "Implemented both frontend and backend modules with a mobile-friendly UI and business-focused usability for daily operations.")

add_section_heading(doc, "Certifications")

cert_1 = doc.add_paragraph(style="List Bullet")
cert_1.paragraph_format.left_indent = Inches(0.18)
cert_1.paragraph_format.space_after = Pt(0)
add_run(cert_1, "Alpha DSA with Java : Apna College (Shadha Khapra)", underline=True)

cert_2 = doc.add_paragraph(style="List Bullet")
cert_2.paragraph_format.left_indent = Inches(0.18)
cert_2.paragraph_format.space_after = Pt(0)
add_run(cert_2, "The Complete Web Developer Course 3.0 : Udemy (31 hours, Rob Percival, Lo destars)", bold=True, underline=True)

doc.save(OUTPUT_PATH)
print(OUTPUT_PATH)
